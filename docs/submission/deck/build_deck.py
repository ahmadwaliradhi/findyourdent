"""Build the AI Builder Cup 2026 pitch deck (issue #25) on the official template.

The template's slides, order and headings are kept; only content is added. The
template's instruction slide ("Note: Please download this template...") is removed.

    python -m venv .venv && .venv/bin/pip install python-pptx 'qrcode[pil]'
    .venv/bin/python docs/submission/deck/build_deck.py

Writes FindYourDent-AI-Builder-Cup-2026.pptx next to this file. Export the PDF with
LibreOffice: soffice --headless --convert-to pdf <pptx>. Figures come from the files
named in each slide's source line; update them there first, then here.
"""

import copy
import io
from pathlib import Path

import qrcode
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).parent
TEMPLATE = HERE / "template.pptx"
OUT = HERE / "FindYourDent-AI-Builder-Cup-2026.pptx"
SNAP = HERE / "snapshots"

FONT = "Google Sans Flex"

# Brand tokens from docs/design/design-system/tokens.json (light theme).
INK = RGBColor(0x17, 0x24, 0x27)
MUTED = RGBColor(0x5B, 0x6B, 0x70)
LINE = RGBColor(0xD5, 0xDD, 0xDF)
SOFT = RGBColor(0xF3, 0xF6, 0xF6)
BRAND = RGBColor(0x0B, 0x6E, 0x6A)
TINT = RGBColor(0xDC, 0xEF, 0xEC)
RED = RGBColor(0xB4, 0x23, 0x18)
RED_TINT = RGBColor(0xFD, 0xEC, 0xEA)
AMBER_TINT = RGBColor(0xFD, 0xF3, 0xD8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

TITLE_Y, TITLE_H = 1.12, 0.62
BODY_Y = 1.9
SOURCE_Y = 5.13

SKI_SOURCE = (
    "Sources: SKI 2023 Dalam Angka, BKPK Kemenkes, Tabel 9.5 p. 324 (56.9%) and Tabel 9.7 p. 327 (57.0%); "
    "BKPK SKI 2023 oral health factsheet, Gambar 3 p. 2 (11.2%)."
)


# ---------- low-level helpers ----------


def _font(run, size, color=INK, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def text(slide, x, y, w, h, paras, size=11, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, space_after=0, line_spacing=None, name=None):
    """paras: str, or list of paragraphs; a paragraph is a str or a list of (text, opts) runs."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        if line_spacing:
            p.line_spacing = line_spacing
        runs = [(para, {})] if isinstance(para, str) else para
        for t, o in runs:
            r = p.add_run()
            r.text = t
            _font(r, o.get("size", size), o.get("color", color), o.get("bold", bold), o.get("italic", False))
    return box


def box(slide, x, y, w, h, fill=SOFT, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08,
        dash=False, line_w=0.75, name=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        s.name = name
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
        if dash:
            s.line.dash_style = 4  # dash
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    style = s._element.find(qn("p:style"))
    if style is not None:  # theme effect refs draw shadows in some renderers
        s._element.remove(style)
    s.text_frame.text = ""
    return s


def label_in(shape, paras, size=10, color=INK, bold=False, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
             margin=0.08):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    tf.vertical_anchor = anchor
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        runs = [(para, {})] if isinstance(para, str) else para
        for t, o in runs:
            r = p.add_run()
            r.text = t
            _font(r, o.get("size", size), o.get("color", color), o.get("bold", bold), o.get("italic", False))


def arrow(slide, x1, y1, x2, y2, color=MUTED, width=1.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    return c


def circle_num(slide, x, y, n, d=0.3, fill=BRAND, color=WHITE, size=10):
    c = box(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    label_in(c, str(n), size=size, color=color, bold=True, margin=0)
    return c


def action_title(slide, t):
    return text(slide, 0.29, TITLE_Y, 9.4, TITLE_H, t, size=19, bold=True, anchor=MSO_ANCHOR.TOP,
                line_spacing=1.0, name="Action title")


def source(slide, t):
    return text(slide, 0.29, SOURCE_Y, 9.4, 0.36, t, size=7.5, color=MUTED, anchor=MSO_ANCHOR.BOTTOM,
                name="Source")


def kicker(slide, x, y, w, t, color=BRAND):
    return text(slide, x, y, w, 0.22, t.upper(), size=8.5, color=color, bold=True)


def takeaway(slide, y, t, h=0.42):
    b = box(slide, 0.29, y, 9.42, h, fill=TINT, radius=0.12)
    label_in(b, [[("So what  ", {"bold": True, "color": BRAND}), (t, {"color": INK})]], size=11,
             align=PP_ALIGN.LEFT, margin=0.18)
    return b


def picture_framed(slide, path, x, y, h):
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), height=Inches(h))
    pic.line.color.rgb = LINE
    pic.line.width = Pt(0.75)
    return pic


def table(slide, x, y, w, rows, col_w, size=9, header=True, row_h=0.26, bold_first_col=False, fills=None):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w),
                                   Inches(row_h * len(rows)))
    tbl = shape.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    style = tblPr.find(qn("a:tableStyleId"))
    if style is not None:
        style.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # No Style, Table Grid
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        tbl.rows[r].height = Inches(row_h)
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.07)
            cell.margin_top = cell.margin_bottom = Inches(0.025)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            is_head = header and r == 0
            fill = (fills or {}).get((r, c))
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill or (TINT if is_head else WHITE)
            opts = val if isinstance(val, tuple) else (val, {})
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = opts[1].get("align", PP_ALIGN.LEFT)
            run = p.add_run()
            run.text = opts[0]
            _font(run, size, opts[1].get("color", BRAND if is_head else INK),
                  is_head or opts[1].get("bold", False) or (bold_first_col and c == 0))
            # thin horizontal rules only
            tcPr = cell._tc.get_or_add_tcPr()
            for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                ln = etree.SubElement(tcPr, qn(tag))
                if tag in ("a:lnB",):
                    ln.set("w", "6350")
                    sf = etree.SubElement(ln, qn("a:solidFill"))
                    clr = etree.SubElement(sf, qn("a:srgbClr"))
                    clr.set("val", "D5DDDF")
                else:
                    ln.set("w", "0")
                    etree.SubElement(ln, qn("a:noFill"))
    return shape


def set_runs_text(paragraph, value, size=None, bold=None):
    """Append a value run after the label run, copying the label's formatting."""
    label = paragraph.runs[0]
    new = copy.deepcopy(label._r)
    label._r.addnext(new)
    run = paragraph.runs[1]
    run.text = value
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    return run


def trim_paragraphs(shape, keep):
    paras = shape.text_frame._txBody.findall(qn("a:p"))
    for p in paras[keep:]:
        p.getparent().remove(p)


def delete_slide(prs, index):
    sldIdLst = prs.slides._sldIdLst
    sld = sldIdLst[index]
    prs.part.drop_rel(sld.rId)
    sldIdLst.remove(sld)


def shape_named(slide, prefix):
    for s in slide.shapes:
        if s.name.startswith(prefix) and s.has_text_frame and s.text_frame.text.strip():
            return s
    raise KeyError(prefix)


def heading_box(slide):
    """The template's heading text box on a content slide (top-left, under the banner)."""
    for s in slide.shapes:
        if s.has_text_frame and s.text_frame.text.strip() and Emu(s.top).inches < 1.0:
            return s
    raise KeyError("heading")


# ---------- slides ----------


def s_team(slide):
    tb = shape_named(slide, "Google Shape;54")
    paras = tb.text_frame.paragraphs
    set_runs_text(paras[2], "FindYourDent")
    set_runs_text(paras[3], "Ahmad Wali Radhi")
    set_runs_text(
        paras[4],
        "Survei Kesehatan Indonesia 2023: 56.9% of Indonesians aged 3 and over had a dental or oral problem "
        "in the past year, yet only 11.2% of them were treated by dental health workers. A Patient cannot tell "
        "how urgent it is, what it will cost, whether BPJS pays, or where to go.",
        size=13,
    )
    trim_paragraphs(tb, 5)  # drop the template's trailing empty paragraph
    tb.top = Inches(2.86)
    text(slide, 0.29, 5.2, 9.4, 0.25, SKI_SOURCE, size=7.5, color=MUTED, name="Source")


def s_brief(slide):
    action_title(slide, "A Patient describes a toothache in their own words and leaves knowing how urgent it is, "
                        "what it costs, and where to go")
    # left: the gap
    kicker(slide, 0.29, BODY_Y, 3.0, "The gap · SKI 2023")
    stats = [
        ("56.9%", "of Indonesians aged 3+ had a dental or oral problem in the past year"),
        ("11.2%", "of them were treated by dental health workers"),
        ("57.0%", "of people not treated by a dental health worker said they felt no need"),
    ]
    for i, (num, lab) in enumerate(stats):
        y = BODY_Y + 0.3 + i * 0.98
        text(slide, 0.29, y, 1.25, 0.5, num, size=26, bold=True, color=BRAND if i != 1 else RED)
        text(slide, 1.6, y + 0.06, 1.85, 0.8, lab, size=10, color=INK)
    # divider
    box(slide, 3.62, BODY_Y, 0.01, 3.15, fill=LINE, shape=MSO_SHAPE.RECTANGLE)
    # right: what we do
    kicker(slide, 3.9, BODY_Y, 5.5, "What FindYourDent does")
    steps = [
        ("Describe", "Type or speak the complaint in Indonesian, a regional language or English. "
                     "The complaint is never stored."),
        ("Guidance", "Urgency tier, likely category, visit journey, cost estimate with why prices vary, and "
                     "BPJS coverage. Red flags go straight to “Ke IGD sekarang”."),
        ("Shortlist", "Nearby Providers that can treat this complaint, with Testimonies and price signals."),
        ("Testimony", "After the visit the Patient tells others. Gemini screens it against fixed rules; "
                      "a proof of visit earns “Kunjungan terbukti”."),
    ]
    for i, (t, d) in enumerate(steps):
        y = BODY_Y + 0.3 + i * 0.73
        circle_num(slide, 3.9, y, i + 1)
        text(slide, 4.35, y - 0.01, 5.3, 0.7,
             [[(t, {"bold": True, "size": 11.5})], [(d, {"size": 10, "color": MUTED})]], line_spacing=1.0)
    source(slide, SKI_SOURCE + " Product: docs/specs/README.md.")


def s_opportunities(slide):
    hb = heading_box(slide)
    prompts = [p.text for p in hb.text_frame.paragraphs][1:4]
    trim_paragraphs(hb, 1)
    hb.height = Inches(0.35)
    action_title(slide, "Others answer one part of the question; FindYourDent turns a complaint into a decision "
                        "and a place Patients can trust")
    xs = [0.29, 3.43, 6.57]
    w, y0, h = 2.98, BODY_Y, 2.62
    for x in xs:
        box(slide, x, y0, w, h, fill=SOFT, radius=0.05)
    for x, ptxt in zip(xs, prompts):
        text(slide, x + 0.15, y0 + 0.12, w - 0.3, 0.42, ptxt, size=10.5, bold=True, color=BRAND)
    # column 1: competitors
    rows = [
        ("Alodokter", "Health articles, doctor chat and booking. Not a structured dental decision with "
                      "cost and BPJS for this complaint."),
        ("Halodoc", "Paid chat with a doctor and medicine delivery. Does not route the complaint to the "
                    "right kind of dental Provider."),
        ("Google Maps", "Finds places with open star ratings. No urgency or cost, and reviews are not "
                        "checked against a visit."),
    ]
    for i, (n, d) in enumerate(rows):
        text(slide, xs[0] + 0.15, y0 + 0.56 + i * 0.68, w - 0.3, 0.66,
             [[(n, {"bold": True, "size": 10})], [(d, {"size": 9, "color": MUTED})]], line_spacing=1.0)
    # column 2: how it solves
    solves = [
        ("Urgency tier", " answers “do I need to go?” for the 57.0% who felt no need."),
        ("Cost range and BPJS door", " before the visit, with why prices vary."),
        ("Red-flag check", " sends emergencies to an IGD before any AI call."),
        ("Shortlist", " of nearby Providers that can treat this complaint."),
    ]
    for i, (b, d) in enumerate(solves):
        text(slide, xs[1] + 0.15, y0 + 0.6 + i * 0.52, w - 0.3, 0.5,
             [[(b, {"bold": True}), (d, {"color": INK})]], size=9.5, line_spacing=1.0)
    # column 3: USP
    text(slide, xs[2] + 0.15, y0 + 0.55, w - 0.3, 1.3,
         [[("“FindYourDent turns a toothache into a guided decision. A grounded Gemini agent explains "
            "urgency, the visit journey, cost and BPJS coverage, then shortlists Providers ranked by Patient "
            "Testimonies that FindYourDent owns, each one screened by AI and checked against proof of "
            "visit.”", {"italic": True})]], size=9.5, line_spacing=1.0)
    for i, tag in enumerate(["Testimonies FindYourDent owns", "Checked against proof of visit",
                             "Ranking never paid"]):
        ty = y0 + 1.86 + i * 0.25
        box(slide, xs[2] + 0.17, ty + 0.07, 0.08, 0.08, fill=BRAND, shape=MSO_SHAPE.OVAL)
        text(slide, xs[2] + 0.34, ty, w - 0.5, 0.25, tag, size=9.5, bold=True, color=BRAND)
    takeaway(slide, 4.66, "The moat is the Testimony data: first-hand, moderated, tied to a visit, and never for sale.")
    source(slide, "Differentiator: docs/specs/README.md. Competitor scope from public product pages; "
                  "parity risk analysis in docs/hackathon-plan.md 1.1.")


def s_features(slide):
    action_title(slide, "Eight features cover the Patient journey end to end, and the Provider side starts with "
                        "a claimed page")
    feats = [
        ("Guidance card", "Urgency, likely category, visit journey, procedures, cost estimate, BPJS and sources."),
        ("Emergency Escalation", "Deterministic red-flag check before Gemini; red “Ke IGD sekarang” card "
                                 "and the nearest IGD."),
        ("Voice input", "Speak the complaint for up to a minute, in Indonesian, a regional language or English."),
        ("Shortlist", "Nearby Providers that can treat the complaint, with reasons, Testimonies and price signals."),
        ("Provider page", "Facts, Testimonies and ratings, WhatsApp and phone; Google rating loads on tap."),
        ("Moderation", "Gemini checks fixed rules; fix-before-submit for phone numbers. Humans "
                                  "remove, AI never deletes."),
        ("Kunjungan terbukti", "A proof of visit marks a Testimony as a proven visit, the defence against buzzers."),
        ("Claim and Price List", "Owner verifies with a SIP; Gemini reads the price list into rows the owner "
                                 "confirms."),
    ]
    w, h, gx, gy = 2.2, 1.3, 0.207, 0.14
    for i, (t, d) in enumerate(feats):
        x = 0.29 + (i % 4) * (w + gx)
        y = BODY_Y + (i // 4) * (h + gy)
        box(slide, x, y, w, h, fill=SOFT, radius=0.06)
        circle_num(slide, x + 0.14, y + 0.14, i + 1, d=0.28, size=9)
        text(slide, x + 0.52, y + 0.16, w - 0.62, 0.3, t, size=10.5, bold=True)
        text(slide, x + 0.14, y + 0.52, w - 0.28, 0.82, d, size=9, color=MUTED, line_spacing=1.0)
    text(slide, 0.29, 4.82, 9.4, 0.25,
         [[("Also live: ", {"bold": True, "color": BRAND}),
           ("shareable Guidance link without the complaint · Indonesian and English · live impact counter "
            "· directory by kota/kabupaten · sign-in only at “Kirim”", {"color": INK})]], size=9.5)
    source(slide, "Specs 01 to 05 and 09, docs/specs/. Every feature listed is in the deployed build at findyourdent.com.")


def s_process(slide):
    action_title(slide, "Safety runs first: a deterministic red-flag check decides before Gemini sees the complaint")
    lanes = [
        ("Guidance", 1.95, ["Complaint, typed or spoken", "Red-flag check (deterministic)",
                            "Gemini Guidance: validated JSON", "Guidance card + knowledge files",
                            "Shortlist and Provider page"]),
        ("Testimony", 3.42, ["Visit the Provider", "Write a Testimony, sign in at Kirim",
                             "Gemini Moderation: which rules fired", "Code decides: publish, fix or hold",
                             "Admin queue: humans can remove"]),
        ("Provider", 4.32, ["Claim the page: SIP + phone check", "Upload price list (PDF or photo)",
                            "Gemini extracts rows", "Owner checks and confirms",
                            "Dated prices on page and Shortlist"]),
    ]
    bw, gap, x0, bh = 1.47, 0.24, 1.25, 0.6
    for name, y, steps in lanes:
        text(slide, 0.29, y, 0.9, bh, name, size=9.5, bold=True, color=BRAND, anchor=MSO_ANCHOR.MIDDLE)
        for i, st in enumerate(steps):
            x = x0 + i * (bw + gap)
            ai = "Gemini" in st
            safety = "Red-flag" in st
            b = box(slide, x, y, bw, bh, fill=TINT if ai else (AMBER_TINT if safety else SOFT),
                    line=BRAND if ai else None, radius=0.1)
            label_in(b, st, size=8.5, bold=ai or safety)
            if i:
                arrow(slide, x - gap + 0.02, y + bh / 2, x - 0.02, y + bh / 2)
    # emergency branch under the red-flag box
    rx = x0 + (bw + gap)
    eb = box(slide, rx - 0.3, 2.86, bw + 0.6, 0.4, fill=RED_TINT, line=RED, radius=0.15)
    label_in(eb, [[("Red flag: ", {"bold": True, "color": RED}), ("“Ke IGD sekarang” + nearest IGD",
                                                                  {"color": RED})]], size=8.5)
    arrow(slide, rx + bw / 2, 1.95 + bh + 0.01, rx + bw / 2, 2.85, color=RED)
    text(slide, rx + bw / 2 + 0.06, 2.6, 0.9, 0.2, "yes", size=8, color=RED, bold=True)
    text(slide, rx + bw + 0.02, 1.95 + 0.06, 0.25, 0.2, "no", size=8, color=MUTED, bold=True)
    # legend
    for i, (fill, line, t) in enumerate([(TINT, BRAND, "Gemini call"), (AMBER_TINT, None, "Deterministic code")]):
        lx = 5.6 + i * 1.6
        box(slide, lx, 3.02, 0.18, 0.14, fill=fill, line=line, radius=0.2)
        text(slide, lx + 0.25, 2.98, 1.3, 0.22, t, size=8, color=MUTED)
    source(slide, "Specs 01, 03 and 05 in docs/specs/; ADR-0004 (AI publishes, humans remove). The model can raise "
                  "urgency but never lower it below a red flag.")


def phone(slide, x, y, w=1.55, h=2.82):
    box(slide, x, y, w, h, fill=WHITE, line=MUTED, radius=0.12, line_w=1.25)
    box(slide, x + 0.55, y + 0.08, w - 1.1, 0.05, fill=LINE, radius=0.5)
    return x + 0.12, y + 0.22, w - 0.24


def bar(slide, x, y, w, h=0.07, fill=LINE):
    box(slide, x, y, w, h, fill=fill, radius=0.5)


def s_wireframes(slide):
    action_title(slide, "Mobile-first: one screen asks, one answers, one shows where to go")
    xs = [0.75, 3.95, 7.15]
    y = BODY_Y - 0.05
    # 1 Home
    ix, iy, iw = phone(slide, xs[0], y)
    bar(slide, ix, iy + 0.05, 0.6, fill=BRAND)
    bar(slide, ix, iy + 0.3, iw, h=0.12, fill=INK)
    bar(slide, ix, iy + 0.48, iw * 0.7, h=0.12, fill=INK)
    for k in range(3):
        bar(slide, ix, iy + 0.72 + k * 0.12, iw * (0.95 - 0.1 * k))
    box(slide, ix, iy + 1.15, iw, 0.62, fill=WHITE, line=LINE, radius=0.08)
    text(slide, ix + 0.06, iy + 1.2, iw - 0.1, 0.3, "Ceritakan keluhan gigi Anda", size=6.5, color=MUTED)
    b = box(slide, ix, iy + 1.88, iw, 0.24, fill=BRAND, radius=0.3)
    label_in(b, "Lihat panduan", size=7, color=WHITE, bold=True)
    b = box(slide, ix, iy + 2.18, iw, 0.24, fill=WHITE, line=BRAND, radius=0.3)
    label_in(b, "Rekam pesan suara", size=7, color=BRAND, bold=True)
    # 2 Guidance
    ix, iy, iw = phone(slide, xs[1], y)
    bar(slide, ix, iy + 0.05, iw * 0.75, h=0.1, fill=INK)
    u = box(slide, ix, iy + 0.25, iw, 0.38, fill=AMBER_TINT, radius=0.1)
    label_in(u, "Segera: 1 to 3 days", size=7, bold=True, align=PP_ALIGN.LEFT)
    for k, c in enumerate([LINE, LINE, RGBColor(0xD9, 0x8E, 0x04), LINE]):
        bar(slide, ix + k * (iw / 4) + 0.01, iy + 0.7, iw / 4 - 0.04, h=0.05, fill=c)
    text(slide, ix, iy + 0.82, iw, 0.18, "Likely problem", size=6.5, color=BRAND, bold=True)
    bar(slide, ix, iy + 1.02, iw * 0.9)
    bar(slide, ix, iy + 1.14, iw * 0.7)
    text(slide, ix, iy + 1.28, iw, 0.18, "Cost estimate", size=6.5, color=BRAND, bold=True)
    for k in range(2):
        r = box(slide, ix, iy + 1.48 + k * 0.3, iw, 0.24, fill=SOFT, radius=0.1)
    text(slide, ix, iy + 2.08, iw, 0.18, "BPJS Kesehatan", size=6.5, color=BRAND, bold=True)
    box(slide, ix, iy + 2.27, iw, 0.28, fill=RGBColor(0xE6, 0xF0, 0xFA), radius=0.1)
    # 3 Shortlist
    ix, iy, iw = phone(slide, xs[2], y)
    bar(slide, ix, iy + 0.05, iw * 0.8, h=0.1, fill=INK)
    for k in range(3):
        cy = iy + 0.28 + k * 0.76
        box(slide, ix, cy, iw, 0.68, fill=SOFT, radius=0.08)
        bar(slide, ix + 0.07, cy + 0.08, iw * 0.6, h=0.08, fill=INK)
        bar(slide, ix + 0.07, cy + 0.22, iw * 0.8)
        b1 = box(slide, ix + 0.07, cy + 0.4, 0.5, 0.18, fill=BRAND, radius=0.3)
        label_in(b1, "Lihat", size=6, color=WHITE, bold=True, margin=0)
        b2 = box(slide, ix + 0.62, cy + 0.4, 0.6, 0.18, fill=WHITE, line=BRAND, radius=0.3)
        label_in(b2, "Maps", size=6, color=BRAND, bold=True, margin=0)
    for i in range(2):
        arrow(slide, xs[i] + 1.65, y + 1.4, xs[i + 1] - 0.1, y + 1.4, color=BRAND, width=1.5)
    caps = [("Ask", "Free text or voice; disclaimer and no-storage notice up front."),
            ("Answer", "Urgency first, then problem, visits, cost and BPJS, with sources."),
            ("Go", "Shortlist cards end in two actions: the Provider page or Google Maps.")]
    for x, (t, d) in zip(xs, caps):
        text(slide, x - 0.3, 4.68, 2.15, 0.42, [[(t + ". ", {"bold": True}), (d, {"color": MUTED})]], size=8.5,
             line_spacing=1.0)
    source(slide, "Wireframes redrawn from the mockups in docs/design/mockups/html (Beranda, Panduan, Darurat, Tempat), "
                  "built mobile-first at 360 px.")


def s_architecture(slide):
    action_title(slide, "Two Cloud Run services in Jakarta; only the private agent calls Gemini")
    def node(x, y, w, h, title, sub, fill=SOFT, line=None, title_color=INK):
        b = box(slide, x, y, w, h, fill=fill, line=line, radius=0.08)
        label_in(b, [[(title, {"bold": True, "color": title_color})], [(sub, {"size": 7.5, "color": MUTED})]],
                 size=9, align=PP_ALIGN.LEFT, margin=0.09)
        return b
    # client
    node(0.29, 2.75, 1.35, 0.95, "Patient or Provider", "Phone browser, Indonesian or English")
    # edge
    node(1.95, 1.85, 1.95, 0.72, "Firebase Hosting", "findyourdent.com, rewrites every path to fyd-web")
    node(1.95, 2.72, 1.95, 0.55, "Firebase Auth", "Google sign-in, only at “Kirim”")
    node(1.95, 3.42, 1.95, 0.62, "App Check", "reCAPTCHA Enterprise guards the web API")
    # Jakarta group
    g = box(slide, 4.15, 1.72, 2.3, 3.36, fill=None, line=BRAND, dash=True, radius=0.04)
    text(slide, 4.27, 1.78, 2.1, 0.2, "asia-southeast2 · Jakarta", size=7.5, color=BRAND, bold=True)
    node(4.27, 2.02, 2.06, 0.66, "Cloud Run fyd-web", "Public. React Router SSR and web API", fill=TINT)
    node(4.27, 2.86, 2.06, 0.66, "Cloud Run fyd-agent", "Private, ID token only. Google ADK + FastAPI",
         fill=TINT)
    node(4.27, 3.7, 2.06, 0.68, "Firestore", "Written by fyd-web only. Never a complaint's text")
    node(4.27, 4.48, 2.06, 0.5, "Cloud Storage", "Proof of visit, price-list uploads")
    # external / AI
    node(6.75, 1.85, 2.95, 0.88, "Gemini on Vertex AI (global)",
         "3.8 Flash: Guidance and Shortlist reasons. 3.5 Flash-Lite: Moderation and price-list extraction",
         fill=TINT, line=BRAND, title_color=BRAND)
    node(6.75, 2.86, 2.95, 0.55, "Agent Search", "Grounding over guideline and BPJS documents; curated files as fallback")
    node(6.75, 3.53, 2.95, 0.55, "Places API (New)", "Provider search and details, under daily caps")
    node(6.75, 4.2, 2.95, 0.78, "Operations", "Secret Manager · Cloud Scheduler + Workflows (warm agent) · "
                                              "Cloud Run Job · Monitoring uptime + alerts · Budget")
    # arrows
    arrow(slide, 1.65, 3.05, 1.94, 2.25)
    arrow(slide, 3.91, 2.21, 4.26, 2.35)
    arrow(slide, 5.3, 2.69, 5.3, 2.85)
    arrow(slide, 6.34, 3.1, 6.74, 2.4, color=BRAND, width=1.25)
    arrow(slide, 6.34, 3.18, 6.74, 3.14)
    arrow(slide, 6.34, 2.45, 6.74, 3.75)
    source(slide, "Deployed build, 18 October 2026: README.md, infra/gcp/README.md, ADR-0003 (Gemini global endpoint), "
                  "ADR-0008 (Hosting in front of Cloud Run). The browser never calls the agent, Places or Gemini.")


def s_tech(slide):
    action_title(slide, "Firebase at the edge, Firestore for data, Cloud Run for compute, Gemini for judgement")
    cards = [
        ("Firebase", "Hosting serves findyourdent.com and rewrites to Cloud Run. Auth signs Patients and Providers "
                     "in with Google. App Check with reCAPTCHA Enterprise guards the web API."),
        ("Firestore", "In Jakarta: Providers, Testimonies, Moderation decisions and owner-confirmed price lists. "
                      "Never a Patient’s complaint text."),
        ("Cloud Run", "In Jakarta: the public web service, and a private Google ADK agent service that only the "
                      "web service can call. One warm agent instance during judging."),
        ("Gemini", "On Vertex AI: writes the Guidance card as validated structured output, moderates every "
                   "Testimony before it is published, reads price lists into rows the owner confirms."),
    ]
    w, gx = 2.2, 0.207
    for i, (t, d) in enumerate(cards):
        x = 0.29 + i * (w + gx)
        box(slide, x, BODY_Y, w, 2.05, fill=TINT if i == 3 else SOFT, radius=0.06)
        text(slide, x + 0.15, BODY_Y + 0.15, w - 0.3, 0.3, t, size=14, bold=True, color=BRAND)
        text(slide, x + 0.15, BODY_Y + 0.55, w - 0.3, 1.45, d, size=9.5, line_spacing=1.0)
    rows = [
        ("More Google Cloud", "Google ADK · Agent Search · Places API (New) · Cloud Storage · "
                              "Secret Manager · Cloud Scheduler · Workflows · Cloud Monitoring"),
        ("Web", "React Router v7 (SSR) · TypeScript · Tailwind v4 · shadcn/ui · i18next"),
        ("Agent", "Python 3.13 · FastAPI · Pydantic v2 · uv"),
        ("Contracts and tests", "JSON Schema (Ajv + generated Pydantic) · Vitest on Firebase emulators "
                                "· pytest · Playwright · Biome · Ruff · GitHub Actions"),
    ]
    for i, (k, v) in enumerate(rows):
        y = 4.05 + i * 0.26
        text(slide, 0.29, y, 1.6, 0.24, k, size=9, bold=True, color=BRAND)
        text(slide, 1.9, y, 7.8, 0.24, v, size=9)
    source(slide, "docs/submission/brief-description.md (issue #109), README.md.")


def s_cost(slide):
    action_title(slide, "Judging traffic costs about USD 17 for 23 days, and daily quotas cap every paid API")
    kicker(slide, 0.29, BODY_Y, 4.9, "USD, 300 sessions, 16 Oct to 7 Nov 2026 (column B2)")
    cd = CategoryChartData()
    cats = ["Warm agent instance (20 days)", "Gemini: Guidance sessions", "Gemini: Moderation, extraction, cache",
            "Cloud Run active time", "Storage and registry", "PPN 11%"]
    vals = [9.07, 3.51, 1.94, 0.32, 0.25, 1.66]
    cd.categories = list(reversed(cats))
    cd.add_series("USD", list(reversed(vals)))
    gf = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.2), Inches(BODY_Y + 0.22), Inches(5.0),
                                Inches(2.55), cd)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    ch.font.name = FONT
    ch.font.size = Pt(8.5)
    ch.font.color.rgb = INK
    plot = ch.plots[0]
    plot.gap_width = 55
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = "0.00"
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(8.5)
    dl.font.color.rgb = INK
    ser = plot.series[0]
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = BRAND
    pt = ser.points[len(vals) - 1]  # top bar (warm instance) after reverse
    va = ch.value_axis
    va.visible = False
    va.has_major_gridlines = False
    va.maximum_scale = 11
    ca = ch.category_axis
    ca.format.line.fill.background()
    ca.tick_labels.font.size = Pt(8.5)
    ca.tick_labels.font.color.rgb = MUTED
    del pt
    text(slide, 0.29, 4.62, 4.9, 0.45,
         [[("Total USD 16.75 with PPN. ", {"bold": True}),
           ("At 2,000 sessions: USD 37.77. The warm instance is the biggest line; Cloud Run’s free tier should "
            "cut it.", {"color": MUTED})]], size=9, line_spacing=1.0)
    # quotas
    kicker(slide, 5.55, BODY_Y, 4.1, "Hard caps set in the project")
    rows = [
        ["Quota", "Daily cap"],
        ["Places Text Search", "160 calls"],
        ["Places Get Place (details, reviews)", "83 calls"],
        ["Places Autocomplete", "300 calls"],
        ["Places Nearby, Photos", "0"],
        ["Gemini 3.8 Flash input tokens", "2,000,000"],
        ["Gemini 3.5 Flash-Lite input tokens", "1,000,000"],
        ["Billing budget, whole account", "IDR 1.65 M (~USD 100), alerts 50% and 90%"],
    ]
    table(slide, 5.55, BODY_Y + 0.27, 4.15, rows, [2.25, 1.9], size=8.5, row_h=0.33)
    source(slide, "docs/research/cost-comparison-2026-10-17.md (list prices read 16 to 17 Oct 2026, ex-tax plus PPN 11%; "
                  "no credits); infra/gcp/README.md daily caps. Places spend above free allowances capped at USD 10 a "
                  "month (ADR-0005).")


def s_snapshots(slide):
    action_title(slide, "The live build at findyourdent.com, captured at phone width")
    shots = [("home", "Home: describe the complaint"), ("guidance", "Guidance: urgency, problem, visits"),
             ("emergency", "Emergency: “Ke IGD sekarang”"), ("shortlist", "Shortlist: who can treat it nearby"),
             ("provider", "Provider page: Testimonies, ratings")]
    h = 2.95
    w = h * 390 / 844
    gap = (9.42 - 5 * w) / 4
    for i, (f, cap) in enumerate(shots):
        x = 0.29 + i * (w + gap)
        picture_framed(slide, SNAP / f"{f}.png", x, BODY_Y - 0.12, h)
        text(slide, x - 0.12, BODY_Y + h - 0.06, w + 0.24, 0.3, cap, size=8.5, align=PP_ALIGN.CENTER, color=INK)
    source(slide, "Screenshots of https://findyourdent.com at 390 px (Chrome DevTools, mobile emulation), synthetic "
                  "complaints. Testimonies shown are real Patient posts published through Moderation.")


def s_performance(slide):
    action_title(slide, "Every safety gate passes on the live build; urgency-tier accuracy is the one gate still failing")
    tiles = [("48/48", "must-escalate phrases return “Ke IGD sekarang” without calling Gemini", BRAND),
             ("28/28", "ordinary complaints do not escalate (negations, slang, regional languages)", BRAND),
             ("0", "unsafe medicine advice across the 40-case Guidance evaluation", BRAND),
             ("0", "rule-breaking Testimonies published across 30 cases × 3 runs", BRAND)]
    w, h = 2.08, 1.3
    for i, (n, d, c) in enumerate(tiles):
        x = 0.29 + (i % 2) * (w + 0.15)
        y = BODY_Y + (i // 2) * (h + 0.15)
        box(slide, x, y, w, h, fill=SOFT, radius=0.06)
        text(slide, x + 0.14, y + 0.1, w - 0.28, 0.45, n, size=24, bold=True, color=c)
        text(slide, x + 0.14, y + 0.6, w - 0.28, 0.68, d, size=8.5, color=MUTED, line_spacing=1.0)
    PASS = {"color": BRAND, "bold": True}
    FAIL = {"color": RED, "bold": True}
    rows = [
        ["Metric (live models, 18 Oct 2026)", "Result", "Gate"],
        ["Emergency recall, Guidance", "100% (8/8)", ("pass", PASS)],
        ["Urgency-tier accuracy", "76% (28/37)", ("fail, ≥ 85%", FAIL)],
        ["Category / journey accuracy", "100% / 100%", ""],
        ["BPJS accuracy", "87% (20/23)", ""],
        ["Moderation decision accuracy", "96.7% (29/30)", ("pass, ≥ 90%", PASS)],
        ["Moderation neutrality mismatches", "0 pairs", ("pass", PASS)],
        ["Lighthouse mobile: a11y, practices, SEO", "100 / 100 / 100", ""],
        ["Lighthouse performance, e2e tests", ("in progress (T19)", {"color": MUTED}), ""],
    ]
    table(slide, 4.75, BODY_Y, 4.95, rows, [2.6, 1.3, 1.05], size=8.5, row_h=0.29)
    text(slide, 0.29, 4.8, 9.4, 0.35,
         [[("Honest limits: ", {"bold": True}),
           ("8 of the 9 urgency misses are more urgent than the label, the safe direction. Guidance labels are "
            "source-derived and agent-labelled, not yet reviewed by a licensed dentist.", {"color": MUTED})]],
         size=8.5, line_spacing=1.0)
    source(slide, "README.md: Emergency Escalation results, Guidance evaluation (gemini-3.8-flash, evalset 2026-10-17.1), "
                  "Moderation evaluation (gemini-3.5-flash-lite). Lighthouse via Chrome DevTools, 18 Oct 2026.")


def s_future(slide):
    action_title(slide, "Providers and clinics will pay for tools, never for rank")
    cols = [
        ("Revenue from Providers and clinics", [
            "Verified-profile subscription with price-list hosting for claimed pages",
            "Clinic intake: the shared Guidance link becomes the pre-visit note",
            "Payment Method on Testimonies (spec 07) for price signals",
        ], SOFT),
        ("Rules that do not bend", [
            "Ranking is never paid; a claim does not change the Shortlist order",
            "No ads, promotion, discounts or per-lead fees, in line with KODEKGI and MKEK Fatwa 29/2021",
            "Providers can reply or report, never delete a Testimony",
            "Education and navigation, not diagnosis or prescription",
        ], TINT),
        ("Next build", [
            "Lift urgency-tier accuracy past the 85% gate",
            "A licensed dentist reviews the 40-case evalset",
            "Proven Testimonies part 2 and buzzer defence (spec 09)",
            "Kecamatan coverage and more cities",
        ], SOFT),
    ]
    w = 2.98
    for i, (t, items, fill) in enumerate(cols):
        x = 0.29 + i * (w + 0.24)
        box(slide, x, BODY_Y, w, 2.85, fill=fill, radius=0.05)
        text(slide, x + 0.15, BODY_Y + 0.13, w - 0.3, 0.3, t, size=11, bold=True, color=BRAND)
        for k, it in enumerate(items):
            y = BODY_Y + 0.55 + k * 0.56
            box(slide, x + 0.16, y + 0.06, 0.08, 0.08, fill=BRAND, shape=MSO_SHAPE.OVAL)
            text(slide, x + 0.34, y, w - 0.5, 0.55, it, size=9.5, line_spacing=1.0)
    takeaway(slide, 4.75, "Trust is the product: if rank could be bought, the Testimonies would stop meaning anything.",
             h=0.36)
    source(slide, "Specs 05, 07 and 09 in docs/specs/; monetisation and ethics constraints in docs/hackathon-plan.md 1.2; "
                  "Terms and About pages on findyourdent.com.")


def qr_png(url):
    img = qrcode.make(url, border=1)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return buf


def s_links(slide):
    hb = heading_box(slide)
    paras = hb.text_frame.paragraphs
    links = {
        "GitHub Public Repository": "https://github.com/ahmadwaliradhi/findyourdent",
        "Demo Video Link (3 Minutes)": "https://youtu.be/KsSCkUG_M00",
        "Final Product Link": "https://findyourdent.com",
    }
    for p in paras:
        if p.text.strip() in links:
            url = links[p.text.strip()]
            r = set_runs_text(p, ": ")
            r._r.addnext(copy.deepcopy(r._r))
            u = p.runs[2]
            u.text = url
            u.font.color.rgb = BRAND
            u.hyperlink.address = url
    cards = [("Code", "github.com/ahmadwaliradhi/findyourdent", "https://github.com/ahmadwaliradhi/findyourdent"),
             ("Demo video", "youtu.be/KsSCkUG_M00", "https://youtu.be/KsSCkUG_M00"),
             ("Live product", "findyourdent.com", "https://findyourdent.com")]
    w = 2.98
    for i, (t, d, url) in enumerate(cards):
        x = 0.29 + i * (w + 0.24)
        y = 2.35
        box(slide, x, y, w, 2.45, fill=SOFT, radius=0.05)
        text(slide, x + 0.18, y + 0.15, w - 0.36, 0.3, t, size=12, bold=True, color=BRAND)
        link = text(slide, x + 0.18, y + 0.47, w - 0.36, 0.4, d, size=9.5, color=INK)
        link.text_frame.paragraphs[0].runs[0].hyperlink.address = url
        slide.shapes.add_picture(qr_png(url), Inches(x + 0.18), Inches(y + 0.92), height=Inches(1.35))
    source(slide, "Each link opened in a logged-out browser before submission. The live demo stays up through the "
                  "finalist announcement on 7 November 2026.")


def s_sources(slide):
    text(slide, 0.39, 0.83, 9.3, 0.35, "Sources and notes", size=16, name="Heading")
    rows = [
        ("SKI 2023 figures", "Survei Kesehatan Indonesia 2023 Dalam Angka, BKPK Kemenkes, Bab 9: Tabel 9.5 p. 324, "
                             "Tabel 9.7 p. 327. badankebijakan.kemkes.go.id/ski-2023-dalam-angka"),
        ("11.2% treated", "BKPK SKI 2023 oral health factsheet, Gambar 3 p. 2. "
                          "repository.badankebijakan.kemkes.go.id/5534"),
        ("Reading the 57.0%", "Its denominator is everyone aged 3+ not treated by a dental health worker, including "
                              "people with no dental problem (multiple response)."),
        ("Cost", "docs/research/cost-comparison-2026-10-17.md, USD list prices, ex-tax, PPN 11% as one line, no credits."),
        ("Evaluations", "README.md results sections; evalsets in services/agent/evals/. Guidance labels agent-labelled, "
                        "Moderation labels team-labelled."),
        ("Ethics", "KODEKGI (Kode Etik Kedokteran Gigi Indonesia) and MKEK Fatwa 29/2021 on advertising by doctors "
                   "and dentists."),
    ]
    for i, (k, v) in enumerate(rows):
        y = 1.3 + i * 0.6
        text(slide, 0.29, y, 2.0, 0.5, k, size=10, bold=True, color=BRAND)
        text(slide, 2.35, y, 7.35, 0.55, v, size=9.5, line_spacing=1.0)


def main():
    prs = Presentation(str(TEMPLATE))
    delete_slide(prs, 1)  # template instruction note
    s = list(prs.slides)
    builders = [s_team, s_brief, s_opportunities, s_features, s_process, s_wireframes, s_architecture, s_tech,
                s_cost, s_snapshots, s_performance, s_future, s_links, s_sources]
    for slide, build in zip(s, builders):
        build(slide)
    for slide in prs.slides:  # product owner: drop "(if any)" and "(optional)" from template headings
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    for r in p.runs:
                        for tag in ("(if any)", "(optional)"):
                            if tag in r.text:
                                r.text = r.text.replace(" " + tag, "").replace(tag, "")
    prs.save(str(OUT))
    print(f"wrote {OUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
